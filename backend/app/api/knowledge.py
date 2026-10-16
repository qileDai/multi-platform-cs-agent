"""知识库接口：FAQ CRUD、文档上传、召回测试台、未命中问题、违禁词库。"""
import io
import json
import logging
from datetime import datetime

from fastapi import APIRouter, Depends, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from ..config import settings
from ..core import audit
from ..database import get_db
from ..models import Agent, BannedWord, KnowledgeChunk, KnowledgeDoc, MissedQuestion
from ..rag import ingest, pipeline
from ..rag.ingest import load_similar_questions
from ..schemas import (BannedWordIn, BannedWordOut, DialogImportOut, DocBatchIn, DocBatchOut,
                       DocMetaUpdate, DocStatusUpdate, FaqCreate, FaqPackOut, FaqUpdate,
                       KnowledgeChunkOut, KnowledgeDocOut, MissedQuestionOut,
                       RecallTestRequest, RecallTestResult)
from .deps import get_current_agent, require_admin

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/knowledge", tags=["knowledge"])


def _parse_faq(content: str) -> tuple[str, str]:
    first, _, rest = (content or "").partition("\n")
    question = first.removeprefix("问：")
    answer = rest.removeprefix("答：")
    return question, answer


def _dump_similar(items: list[str] | None) -> str:
    cleaned = [s.strip() for s in (items or []) if s and s.strip()]
    return json.dumps(cleaned, ensure_ascii=False)


def _doc_out(db: Session, doc: KnowledgeDoc) -> KnowledgeDocOut:
    count = db.query(KnowledgeChunk).filter(KnowledgeChunk.doc_id == doc.id).count()
    question, answer = _parse_faq(doc.content) if doc.doc_type == "faq" else ("", "")
    return KnowledgeDocOut(
        id=doc.id, title=doc.title, doc_type=doc.doc_type, status=doc.status,
        version=doc.version, chunk_count=count, question=question, answer=answer,
        category=doc.category or "未分类",
        similar_questions=load_similar_questions(doc.similar_questions),
        hit_count=doc.hit_count or 0,
        created_at=doc.created_at, updated_at=doc.updated_at,
    )


# ============ FAQ ============

@router.get("/docs", response_model=list[KnowledgeDocOut])
def list_docs(agent: Agent = Depends(get_current_agent), db: Session = Depends(get_db)):
    docs = (
        db.query(KnowledgeDoc)
        .filter(KnowledgeDoc.status != "archived")
        .order_by(KnowledgeDoc.updated_at.desc())
        .all()
    )
    return [_doc_out(db, d) for d in docs]


@router.post("/faq", response_model=KnowledgeDocOut)
async def create_faq(req: FaqCreate, agent: Agent = Depends(get_current_agent),
                     db: Session = Depends(get_db)):
    exists = (
        db.query(KnowledgeDoc)
        .filter(
            KnowledgeDoc.doc_type == "faq",
            KnowledgeDoc.title == req.title,
            KnowledgeDoc.status != "archived",
        )
        .first()
    )
    if exists is not None:
        raise HTTPException(409, "该标题已存在")
    doc = KnowledgeDoc(
        title=req.title, doc_type="faq",
        content=f"问：{req.question}\n答：{req.answer}",
        category=(req.category or "未分类").strip() or "未分类",
        similar_questions=_dump_similar(req.similar_questions),
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    audit.log(db, agent, "kb_faq_create", target=f"doc:{doc.id}", detail=req.title)
    await ingest.ingest_faq(doc.id)
    return _doc_out(db, doc)


@router.put("/faq/{doc_id}", response_model=KnowledgeDocOut)
async def update_faq(doc_id: int, req: FaqUpdate, agent: Agent = Depends(get_current_agent),
                     db: Session = Depends(get_db)):
    doc = db.get(KnowledgeDoc, doc_id)
    if doc is None or doc.doc_type != "faq":
        raise HTTPException(404, "FAQ 不存在")
    # content 格式：第一行「问：xxx」，其余「答：xxx」
    question = req.question if req.question is not None else doc.content.split("\n")[0].removeprefix("问：")
    answer = req.answer if req.answer is not None else doc.content.split("\n", 1)[1].removeprefix("答：") if "\n" in doc.content else ""
    doc.content = f"问：{question}\n答：{answer}"
    if req.title:
        doc.title = req.title
    if req.category is not None:
        doc.category = req.category.strip() or "未分类"
    if req.similar_questions is not None:
        doc.similar_questions = _dump_similar(req.similar_questions)
    doc.version += 1
    doc.updated_at = datetime.utcnow()
    db.commit()
    audit.log(db, agent, "kb_faq_update", target=f"doc:{doc.id}", detail=doc.title)
    await ingest.ingest_faq(doc.id)  # 增量重建索引
    return _doc_out(db, doc)


async def _apply_doc_action(db: Session, ids: list[int], action: str) -> tuple[int, list[int]]:
    """改状态并按新状态对齐索引。已经是目标状态的跳过。"""
    if action not in ("active", "disabled", "delete"):
        raise HTTPException(400, "动作只能是启用、停用或删除")
    unique_ids = list(dict.fromkeys(ids))
    if not unique_ids:
        raise HTTPException(400, "请选择知识")
    target = "archived" if action == "delete" else action
    docs = db.query(KnowledgeDoc).filter(KnowledgeDoc.id.in_(unique_ids)).all()
    by_id = {doc.id: doc for doc in docs}
    failed: list[int] = []
    changed: list[int] = []
    now = datetime.utcnow()
    for doc_id in unique_ids:
        doc = by_id.get(doc_id)
        if doc is None or (action != "delete" and doc.status == "archived"):
            failed.append(doc_id)
            continue
        if doc.status == target:
            continue
        doc.status = target
        doc.updated_at = now
        changed.append(doc_id)
    db.commit()
    if not changed:
        return 0, failed
    updated, index_failed = await ingest.sync_doc_index(changed)
    failed.extend(index_failed)
    return updated, failed


@router.post("/docs/batch", response_model=DocBatchOut)
async def batch_docs(req: DocBatchIn, agent: Agent = Depends(get_current_agent),
                     db: Session = Depends(get_db)):
    updated, failed = await _apply_doc_action(db, req.ids, req.action)
    audit.log(
        db, agent, "kb_doc_batch", target="knowledge",
        detail=f"{req.action} updated={updated} failed={len(failed)}",
    )
    return DocBatchOut(updated=updated, failed_ids=failed)


@router.delete("/docs/{doc_id}")
async def delete_doc(doc_id: int, agent: Agent = Depends(get_current_agent),
                     db: Session = Depends(get_db)):
    doc = db.get(KnowledgeDoc, doc_id)
    if doc is None:
        raise HTTPException(404, "文档不存在")
    title = doc.title
    await _apply_doc_action(db, [doc_id], "delete")
    audit.log(db, agent, "kb_doc_delete", target=f"doc:{doc_id}", detail=title)
    return {"ok": True}


@router.patch("/docs/{doc_id}", response_model=KnowledgeDocOut)
async def update_doc_meta(doc_id: int, req: DocMetaUpdate, agent: Agent = Depends(get_current_agent),
                          db: Session = Depends(get_db)):
    """文档（及 FAQ）的标题、分类。标题变更且仍启用时重建索引，分类不进检索。"""
    doc = db.get(KnowledgeDoc, doc_id)
    if doc is None or doc.status == "archived":
        raise HTTPException(404, "文档不存在")
    title_changed = False
    if req.title is not None and req.title.strip():
        title_changed = req.title.strip() != doc.title
        doc.title = req.title.strip()
    if req.category is not None:
        doc.category = req.category.strip() or "未分类"
    doc.updated_at = datetime.utcnow()
    db.commit()
    audit.log(db, agent, "kb_doc_meta", target=f"doc:{doc.id}", detail=doc.title)
    if title_changed and doc.status == "active":
        if doc.doc_type == "faq":
            await ingest.ingest_faq(doc.id)
        else:
            await ingest.ingest_document(doc.id)
    return _doc_out(db, doc)


@router.patch("/docs/{doc_id}/status", response_model=KnowledgeDocOut)
async def set_doc_status(doc_id: int, req: DocStatusUpdate, agent: Agent = Depends(get_current_agent),
                         db: Session = Depends(get_db)):
    doc = db.get(KnowledgeDoc, doc_id)
    if doc is None or doc.status == "archived":
        raise HTTPException(404, "文档不存在")
    if req.status not in ("active", "disabled"):
        raise HTTPException(400, "状态只能是启用或停用")
    previous = doc.status
    updated, failed = await _apply_doc_action(db, [doc_id], req.status)
    if doc_id in failed:
        raise HTTPException(500, "索引更新失败")
    if previous != req.status and updated:
        audit.log(db, agent, "kb_doc_status", target=f"doc:{doc.id}", detail=req.status)
    db.refresh(doc)
    return _doc_out(db, doc)


@router.get("/docs/{doc_id}/chunks", response_model=list[KnowledgeChunkOut])
def list_chunks(doc_id: int, agent: Agent = Depends(get_current_agent),
                db: Session = Depends(get_db)):
    doc = db.get(KnowledgeDoc, doc_id)
    if doc is None:
        raise HTTPException(404, "文档不存在")
    rows = (
        db.query(KnowledgeChunk)
        .filter(KnowledgeChunk.doc_id == doc_id)
        .order_by(KnowledgeChunk.chunk_index)
        .all()
    )
    return [
        KnowledgeChunkOut(id=c.id, chunk_index=c.chunk_index, parent_id=c.parent_id, content=c.content)
        for c in rows
    ]


# ============ 文档上传 ============

@router.post("/upload", response_model=KnowledgeDocOut)
async def upload_doc(file: UploadFile, agent: Agent = Depends(get_current_agent),
                     db: Session = Depends(get_db)):
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件不能超过 10MB")
    try:
        text = ingest.parse_file(file.filename or "unknown.txt", data)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, f"文档解析失败: {exc}")
    if not text.strip():
        raise HTTPException(400, "文档内容为空")

    doc = KnowledgeDoc(title=file.filename or "未命名文档", doc_type="file", content=text)
    db.add(doc)
    db.commit()
    db.refresh(doc)
    audit.log(db, agent, "kb_doc_upload", target=f"doc:{doc.id}", detail=doc.title)
    await ingest.ingest_document(doc.id)
    return _doc_out(db, doc)


def _similar_lines(text: str) -> list[str]:
    return [line.strip() for line in (text or "").splitlines() if line.strip()]


@router.post("/import-dialogs", response_model=DialogImportOut)
async def import_dialogs(file: UploadFile, agent: Agent = Depends(get_current_agent),
                         db: Session = Depends(get_db)):
    """美洽历史对话表洗成 FAQ。有答案的以停用入库，待补不建文档。"""
    name = (file.filename or "").lower()
    if not name.endswith(".xlsx"):
        raise HTTPException(400, "只接受美洽导出的 xlsx")
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件不能超过 10MB")
    from scripts.clean_meiqia_to_faq import clean_xlsx_bytes
    try:
        faq_rows, gap_rows, stats = clean_xlsx_bytes(data)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc

    created = 0
    updated = 0
    reindex_ids: list[int] = []
    for row in faq_rows:
        title = (row.get("标题") or "").strip()
        question = (row.get("标准问") or "").strip()
        answer = (row.get("答案") or "").strip()
        if not title or not question or not answer:
            continue
        category = (row.get("分类") or "未分类").strip() or "未分类"
        similars = _similar_lines(row.get("相似问") or "")
        content = f"问：{question}\n答：{answer}"
        doc = (
            db.query(KnowledgeDoc)
            .filter(
                KnowledgeDoc.doc_type == "faq",
                KnowledgeDoc.title == title,
                KnowledgeDoc.status != "archived",
            )
            .order_by(KnowledgeDoc.id.asc())
            .first()
        )
        if doc is None:
            doc = KnowledgeDoc(
                title=title, doc_type="faq", content=content, category=category,
                similar_questions=_dump_similar(similars), status="disabled",
            )
            db.add(doc)
            db.flush()
            created += 1
            continue
        doc.content = content
        doc.category = category
        doc.similar_questions = _dump_similar(similars)
        doc.version = (doc.version or 1) + 1
        doc.updated_at = datetime.utcnow()
        updated += 1
        if doc.status == "active":
            reindex_ids.append(doc.id)
    db.commit()
    for doc_id in reindex_ids:
        await ingest.ingest_faq(doc_id)
    audit.log(
        db, agent, "kb_dialog_import", target="knowledge",
        detail=f"created={created} updated={updated} gap={len(gap_rows)}",
    )
    return DialogImportOut(
        imported=created + updated,
        created=created,
        updated=updated,
        skipped_invalid=int(stats["invalid"]),
        skipped_missed=int(stats["missed"]),
        gap_count=len(gap_rows),
        gap_titles=[row.get("标题") or "" for row in gap_rows if row.get("标题")],
    )


FAQ_PACK_HEADERS = ("标题", "分类", "标准问", "相似问", "答案", "状态")
_REQUIRED_PACK_HEADERS = ("标题", "标准问", "答案")


class _PackRow:
    def __init__(self, title: str, category: str, question: str, answer: str,
                 similars: list[str], status: str):
        self.title = title
        self.category = category
        self.question = question
        self.answer = answer
        self.similars = similars
        self.status = status


class _PackWrite:
    def __init__(self, doc: KnowledgeDoc | None, row: _PackRow, content_changed: bool):
        self.doc = doc
        self.row = row
        self.content_changed = content_changed


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _load_pack_rows(data: bytes) -> tuple[list[_PackRow], list[str]]:
    """解析 FAQ 表。文件内重复标题以后出现的行为准，返回 (rows, duplicates)。"""
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, "无法读取表格") from exc
    try:
        sheet = workbook.active
        row_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(row_iter)
        except StopIteration:
            raise HTTPException(400, "表格是空的") from None
        headers = [_cell_text(cell) for cell in (header_row or ())]
        missing = [name for name in _REQUIRED_PACK_HEADERS if name not in headers]
        if missing:
            raise HTTPException(400, "表格缺少列：" + "、".join(missing))
        index = {name: headers.index(name) for name in FAQ_PACK_HEADERS if name in headers}

        def column(raw, name: str) -> str:
            pos = index.get(name)
            if pos is None or raw is None or pos >= len(raw):
                return ""
            return _cell_text(raw[pos])

        rows: list[_PackRow] = []
        errors: list[str] = []
        duplicates: list[str] = []
        seen: dict[str, int] = {}
        for offset, raw in enumerate(row_iter, start=2):
            if raw is None or all(_cell_text(cell) == "" for cell in raw):
                continue
            title = column(raw, "标题")
            question = column(raw, "标准问")
            answer = column(raw, "答案")
            status_text = column(raw, "状态")
            if not title or not question or not answer:
                errors.append(f"第 {offset} 行缺少标题、标准问或答案")
                continue
            if status_text in ("", "停用"):
                status = "disabled"
            elif status_text == "启用":
                status = "active"
            else:
                errors.append(f"第 {offset} 行状态无法识别")
                continue
            row = _PackRow(
                title=title,
                category=column(raw, "分类") or "未分类",
                question=question,
                answer=answer,
                similars=_similar_lines(column(raw, "相似问")),
                status=status,
            )
            pos = seen.get(title)
            if pos is None:
                seen[title] = len(rows)
                rows.append(row)
            else:
                # 重复标题：后面的行覆盖前面的行
                rows[pos] = row
                if title not in duplicates:
                    duplicates.append(title)
        if errors:
            shown = errors[:8]
            if len(errors) > 8:
                shown.append(f"另有 {len(errors) - 8} 个问题")
            raise HTTPException(400, "；".join(shown))
        if not rows:
            raise HTTPException(400, "表格里没有 FAQ")
        return rows, duplicates
    finally:
        workbook.close()


def _pack_plan(db: Session, rows: list[_PackRow]):
    existing = (
        db.query(KnowledgeDoc)
        .filter(KnowledgeDoc.doc_type == "faq", KnowledgeDoc.status != "archived")
        .all()
    )
    by_title: dict[str, list[KnowledgeDoc]] = {}
    for doc in existing:
        by_title.setdefault(doc.title, []).append(doc)
    writes: list[_PackWrite] = []
    skipped: list[str] = []
    created = content_updated = enabling = disabling_active = 0
    for row in rows:
        matches = by_title.get(row.title, [])
        if len(matches) > 1:
            skipped.append(row.title)
            continue
        content = f"问：{row.question}\n答：{row.answer}"
        if not matches:
            created += 1
            if row.status == "active":
                enabling += 1
            writes.append(_PackWrite(None, row, False))
            continue
        doc = matches[0]
        content_changed = (
            doc.content != content
            or (doc.category or "未分类") != row.category
            or load_similar_questions(doc.similar_questions) != row.similars
        )
        if content_changed:
            content_updated += 1
        if row.status == "active" and doc.status != "active":
            enabling += 1
        if doc.status == "active" and row.status != "active":
            disabling_active += 1
        if content_changed or doc.status != row.status:
            writes.append(_PackWrite(doc, row, content_changed))
    skipped.sort()
    return writes, created, content_updated, enabling, disabling_active, skipped


def _apply_pack_writes(db: Session, writes: list[_PackWrite]) -> list[int]:
    index_ids: list[int] = []
    now = datetime.utcnow()
    for item in writes:
        row = item.row
        content = f"问：{row.question}\n答：{row.answer}"
        similars = _dump_similar(row.similars)
        if item.doc is None:
            doc = KnowledgeDoc(
                title=row.title, doc_type="faq", content=content, category=row.category,
                similar_questions=similars, status=row.status,
            )
            db.add(doc)
            db.flush()
            if row.status == "active":
                index_ids.append(doc.id)
            continue
        was_active = item.doc.status == "active"
        item.doc.content = content
        item.doc.category = row.category
        item.doc.similar_questions = similars
        item.doc.status = row.status
        item.doc.version = (item.doc.version or 1) + 1
        item.doc.updated_at = now
        if row.status == "active" or was_active:
            index_ids.append(item.doc.id)
    db.commit()
    return index_ids


def _pack_response(applied: bool, created: int, content_updated: int, enabling: int,
                   disabling_active: int, skipped: list[str],
                   duplicates_overridden: list[str] | None = None) -> FaqPackOut:
    return FaqPackOut(
        applied=applied, created=created, content_updated=content_updated,
        enabling=enabling, disabling_active=disabling_active, skipped_ambiguous=skipped,
        duplicates_overridden=duplicates_overridden or [],
    )


@router.get("/faq/export")
def export_faq_pack(agent: Agent = Depends(get_current_agent), db: Session = Depends(get_db)):
    from openpyxl import Workbook

    docs = (
        db.query(KnowledgeDoc)
        .filter(KnowledgeDoc.doc_type == "faq", KnowledgeDoc.status == "active")
        .order_by(KnowledgeDoc.id.asc())
        .all()
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "FAQ"
    sheet.append(list(FAQ_PACK_HEADERS))
    for doc in docs:
        question, answer = _parse_faq(doc.content)
        sheet.append([
            doc.title,
            doc.category or "未分类",
            question,
            "\n".join(load_similar_questions(doc.similar_questions)),
            answer,
            "启用",
        ])
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="faq.xlsx"'},
    )


@router.post("/faq/import", response_model=FaqPackOut)
async def import_faq_pack(file: UploadFile, apply: str = Form("false"),
                          agent: Agent = Depends(get_current_agent), db: Session = Depends(get_db)):
    """带状态的 FAQ 表。apply 不是 true 时只校验和统计，不写库。"""
    name = (file.filename or "").lower()
    if not name.endswith(".xlsx"):
        raise HTTPException(400, "只接受 xlsx")
    data = await file.read()
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(400, "文件不能超过 10MB")
    rows, duplicates = _load_pack_rows(data)
    writes, created, content_updated, enabling, disabling_active, skipped = _pack_plan(db, rows)
    do_apply = apply.strip().lower() in {"1", "true", "yes", "on"}
    if not do_apply:
        return _pack_response(False, created, content_updated, enabling, disabling_active,
                              skipped, duplicates)
    index_ids = _apply_pack_writes(db, writes)
    if index_ids:
        await ingest.sync_doc_index(index_ids)
    audit.log(
        db, agent, "kb_faq_pack_import", target="knowledge",
        detail=(f"created={created} content_updated={content_updated} enabling={enabling} "
                f"disabling={disabling_active} duplicates={len(duplicates)}"),
    )
    return _pack_response(True, created, content_updated, enabling, disabling_active,
                          skipped, duplicates)


# ============ 召回测试台 ============

@router.post("/recall-test", response_model=RecallTestResult)
async def recall_test(req: RecallTestRequest, agent: Agent = Depends(get_current_agent)):
    result = await pipeline.retrieve(req.query, history=[])
    return RecallTestResult(
        rewritten_queries=result.rewritten_queries,
        hits=result.contexts,
        rerank_top_score=result.rerank_top_score,
        threshold=settings.rag_rerank_threshold,
        passed=result.passed,
        degraded=result.degraded,
    )


# ============ 未命中问题 ============

@router.get("/missed", response_model=list[MissedQuestionOut])
def list_missed(agent: Agent = Depends(get_current_agent), db: Session = Depends(get_db)):
    rows = (
        db.query(MissedQuestion)
        .filter(MissedQuestion.status == "pending")
        .order_by(MissedQuestion.count.desc())
        .limit(100)
        .all()
    )
    return [MissedQuestionOut.model_validate(r) for r in rows]


@router.post("/missed/{missed_id}/resolve")
async def resolve_missed(missed_id: int, req: FaqCreate,
                         agent: Agent = Depends(get_current_agent), db: Session = Depends(get_db)):
    """未命中问题一键转 FAQ。"""
    missed = db.get(MissedQuestion, missed_id)
    if missed is None:
        raise HTTPException(404, "记录不存在")
    exists = (
        db.query(KnowledgeDoc)
        .filter(
            KnowledgeDoc.doc_type == "faq",
            KnowledgeDoc.title == req.title,
            KnowledgeDoc.status != "archived",
        )
        .first()
    )
    if exists is not None:
        raise HTTPException(409, "该标题已存在")
    doc = KnowledgeDoc(
        title=req.title, doc_type="faq",
        content=f"问：{req.question}\n答：{req.answer}",
        category=(req.category or "未分类").strip() or "未分类",
        similar_questions=_dump_similar(req.similar_questions),
    )
    db.add(doc)
    missed.status = "resolved"
    db.commit()
    db.refresh(doc)
    await ingest.ingest_faq(doc.id)
    return {"ok": True, "doc_id": doc.id}


# ============ 违禁词库 ============

@router.get("/banned-words", response_model=list[BannedWordOut])
def list_banned(agent: Agent = Depends(get_current_agent), db: Session = Depends(get_db)):
    return [BannedWordOut.model_validate(w) for w in db.query(BannedWord).all()]


@router.post("/banned-words", response_model=BannedWordOut)
def add_banned(req: BannedWordIn, agent: Agent = Depends(require_admin),
               db: Session = Depends(get_db)):
    word = BannedWord()
    _apply_banned(word, req)
    db.add(word)
    _commit_banned(db)
    db.refresh(word)
    _reload_banned_words(db)
    return word


@router.put("/banned-words/{word_id}", response_model=BannedWordOut)
def update_banned(word_id: int, req: BannedWordIn, agent: Agent = Depends(require_admin),
                  db: Session = Depends(get_db)):
    word = db.get(BannedWord, word_id)
    if word is None:
        raise HTTPException(404, "不存在")
    _apply_banned(word, req)
    _commit_banned(db)
    db.refresh(word)
    _reload_banned_words(db)
    return word


@router.delete("/banned-words/{word_id}")
def delete_banned(word_id: int, agent: Agent = Depends(require_admin),
                  db: Session = Depends(get_db)):
    word = db.get(BannedWord, word_id)
    if word:
        db.delete(word)
        db.commit()
        _reload_banned_words(db)
    return {"ok": True}


def _apply_banned(word: BannedWord, req: BannedWordIn) -> None:
    text = (req.word or "").strip()
    if not text:
        raise HTTPException(400, "违禁词不能为空")
    word.word = text
    word.category = (req.category or "极限词").strip() or "极限词"
    word.direction = req.direction


def _commit_banned(db: Session) -> None:
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "该违禁词已存在")


def _reload_banned_words(db: Session):
    from ..core import contentfilter
    contentfilter.load_db_words(
        [(w.word, w.category, w.direction or "out") for w in db.query(BannedWord).all()])
