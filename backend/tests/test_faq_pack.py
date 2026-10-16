"""FAQ 批量启用停用删除，以及带状态的表格导入导出。"""
import io
import json
import uuid

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api.deps import get_current_agent
from app.core.security import hash_password
from app.main import app
from app.models import Agent, KnowledgeChunk, KnowledgeDoc
from app.rag import ingest


def _agent(db) -> Agent:
    agent = Agent(
        username=f"pack_{uuid.uuid4().hex[:8]}", display_name="表格", role="agent",
        status="offline", password_hash=hash_password("x"),
    )
    db.add(agent)
    db.commit()
    return agent


def _faq(db, title: str, answer: str = "原答案", status: str = "disabled", similars: list[str] | None = None) -> KnowledgeDoc:
    doc = KnowledgeDoc(
        title=title, doc_type="faq", content=f"问：{title}吗\n答：{answer}",
        category="价格", status=status,
        similar_questions=json.dumps(similars or [], ensure_ascii=False),
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


def _xlsx(rows: list[list], headers: list[str] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers or ["标题", "分类", "标准问", "相似问", "答案", "状态"])
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _post_pack(client: TestClient, blob: bytes, apply: str):
    return client.post(
        "/api/knowledge/faq/import",
        files={"file": ("faq.xlsx", blob, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"apply": apply},
    )


def _cleanup(db, titles: list[str], agent_id: int):
    docs = db.query(KnowledgeDoc).filter(KnowledgeDoc.title.in_(titles)).all()
    ids = [doc.id for doc in docs]
    if ids:
        db.query(KnowledgeChunk).filter(KnowledgeChunk.doc_id.in_(ids)).delete(synchronize_session=False)
        db.query(KnowledgeDoc).filter(KnowledgeDoc.id.in_(ids)).delete(synchronize_session=False)
    db.query(Agent).filter(Agent.id == agent_id).delete(synchronize_session=False)
    db.commit()


def test_batch_enable_disable_and_delete(db, monkeypatch):
    mark = uuid.uuid4().hex[:8]
    titles = [f"价格{mark}", f"运费{mark}"]
    agent = _agent(db)
    app.dependency_overrides[get_current_agent] = lambda: agent
    left = _faq(db, titles[0])
    right = _faq(db, titles[1])
    manual = (
        "# 价格说明\n"
        "标准款是九十九元，两件还有折扣。\n"
        "\n"
        "# 运费说明\n"
        "满九十九元包邮，偏远地区除外。\n"
    )
    handbook = KnowledgeDoc(title=f"手册{mark}", doc_type="file", content=manual, status="disabled")
    db.add(handbook)
    db.commit()
    db.refresh(handbook)
    titles.append(handbook.title)
    client = TestClient(app)
    client.get("/api/knowledge/docs")
    calls = {"n": 0}
    original = ingest.rebuild_bm25_from_db

    def wrapped():
        calls["n"] += 1
        original()

    monkeypatch.setattr(ingest, "rebuild_bm25_from_db", wrapped)
    try:
        enabled = client.post("/api/knowledge/docs/batch", json={"ids": [left.id, right.id], "action": "active"})
        assert enabled.status_code == 200, enabled.text
        assert enabled.json()["updated"] == 2
        assert enabled.json()["failed_ids"] == []
        assert calls["n"] == 1
        db.expire_all()
        assert db.query(KnowledgeChunk).filter(KnowledgeChunk.doc_id == left.id).count() == 1
        assert db.query(KnowledgeChunk).filter(KnowledgeChunk.doc_id == right.id).count() == 1

        again = client.post("/api/knowledge/docs/batch", json={"ids": [left.id, right.id], "action": "active"})
        assert again.json()["updated"] == 0
        assert calls["n"] == 1

        mixed = client.post("/api/knowledge/docs/batch", json={"ids": [handbook.id], "action": "active"})
        assert mixed.status_code == 200, mixed.text
        db.expire_all()
        pieces = db.query(KnowledgeChunk).filter(KnowledgeChunk.doc_id == handbook.id).all()
        assert len(pieces) >= 2
        assert all("【" in piece.content for piece in pieces)
        assert all(not piece.content.startswith("问：") for piece in pieces)

        disabled = client.post("/api/knowledge/docs/batch", json={"ids": [left.id, right.id], "action": "disabled"})
        assert disabled.json()["updated"] == 2
        db.expire_all()
        assert db.query(KnowledgeChunk).filter(KnowledgeChunk.doc_id == left.id).count() == 0

        removed = client.post("/api/knowledge/docs/batch", json={"ids": [left.id], "action": "delete"})
        assert removed.json()["updated"] == 1
        db.expire_all()
        assert db.get(KnowledgeDoc, left.id).status == "archived"
        listed = client.get("/api/knowledge/docs")
        assert left.id not in [item["id"] for item in listed.json()]
    finally:
        app.dependency_overrides.pop(get_current_agent, None)
        _cleanup(db, titles, agent.id)


def test_faq_pack_preview_roundtrip_and_guards(db):
    mark = uuid.uuid4().hex[:8]
    title = f"套餐{mark}"
    other = f"旁边{mark}"
    ambiguous = f"同名{mark}"
    fresh = f"新题{mark}"
    titles = [title, other, ambiguous, fresh]
    agent = _agent(db)
    app.dependency_overrides[get_current_agent] = lambda: agent
    doc = _faq(db, title, answer="99元", status="active", similars=["怎么卖", "什么价"])
    kept = _faq(db, other, answer="别动", status="active")
    first = _faq(db, ambiguous, answer="一", status="disabled")
    second = _faq(db, ambiguous, answer="二", status="active")
    client = TestClient(app)
    try:
        preview = _post_pack(client, _xlsx([
            [title, "价格", f"{title}吗", "怎么卖\n什么价", "199元", "启用"],
        ]), "false")
        assert preview.status_code == 200, preview.text
        body = preview.json()
        assert body["applied"] is False
        assert body["content_updated"] == 1
        assert body["created"] == 0
        db.expire_all()
        assert "199元" not in db.get(KnowledgeDoc, doc.id).content

        exported = client.get("/api/knowledge/faq/export")
        assert exported.status_code == 200, exported.text
        sheet = load_workbook(io.BytesIO(exported.content), data_only=True).active
        headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        assert headers[:6] == ["标题", "分类", "标准问", "相似问", "答案", "状态"]
        # 只导出启用的：停用条目（答案「一」）不应出现在导出表里
        rows_all = {raw[0]: raw for raw in sheet.iter_rows(min_row=2, values_only=True)}
        assert title in rows_all
        assert rows_all[ambiguous][4] == "二"
        matched = None
        for raw in sheet.iter_rows(min_row=2, values_only=True):
            if raw[0] == title:
                matched = list(raw[:6])
                break
        assert matched is not None
        assert matched[5] == "启用"
        assert "怎么卖" in str(matched[3])
        current = db.get(KnowledgeDoc, doc.id)
        current.content = f"问：{title}吗\n答：改过"
        current.status = "disabled"
        current.similar_questions = "[]"
        db.commit()
        applied = _post_pack(client, _xlsx([matched]), "true")
        assert applied.status_code == 200, applied.text
        assert applied.json()["applied"] is True
        db.expire_all()
        restored = db.get(KnowledgeDoc, doc.id)
        assert restored.status == "active"
        assert "99元" in restored.content
        assert json.loads(restored.similar_questions) == ["怎么卖", "什么价"]
        assert db.query(KnowledgeChunk).filter(KnowledgeChunk.doc_id == doc.id).count() == 1
        assert db.get(KnowledgeDoc, kept.id).content.endswith("别动")

        dup = f"重复{mark}"
        overridden = _post_pack(client, _xlsx([
            [dup, "价格", "问", "相似", "答", "启用"],
            [dup, "价格", "问2", "相似", "答2", "停用"],
        ]), "true")
        titles.append(dup)
        assert overridden.status_code == 200, overridden.text
        assert overridden.json()["duplicates_overridden"] == [dup]
        db.expire_all()
        dup_docs = db.query(KnowledgeDoc).filter(KnowledgeDoc.title == dup).all()
        assert len(dup_docs) == 1
        assert dup_docs[0].content.endswith("答2")
        assert dup_docs[0].status == "disabled"
        assert db.query(KnowledgeChunk).filter(KnowledgeChunk.doc_id == dup_docs[0].id).count() == 0

        no_status = _post_pack(
            client,
            _xlsx([[fresh, "价格", "多少钱", "", "先别上", ""]], ["标题", "分类", "标准问", "相似问", "答案"]),
            "true",
        )
        assert no_status.status_code == 200, no_status.text
        db.expire_all()
        created = db.query(KnowledgeDoc).filter(KnowledgeDoc.title == fresh).one()
        assert created.status == "disabled"
        assert db.query(KnowledgeChunk).filter(KnowledgeChunk.doc_id == created.id).count() == 0

        skipped = _post_pack(client, _xlsx([
            [ambiguous, "价格", "同名吗", "", "新答案", "停用"],
            [f"单独{mark}", "价格", "单独吗", "", "可以", "启用"],
        ]), "true")
        titles.append(f"单独{mark}")
        assert skipped.status_code == 200, skipped.text
        assert ambiguous in skipped.json()["skipped_ambiguous"]
        db.expire_all()
        assert db.get(KnowledgeDoc, first.id).content.endswith("一")
        assert db.get(KnowledgeDoc, second.id).content.endswith("二")
        assert db.get(KnowledgeDoc, second.id).status == "active"
        alone = db.query(KnowledgeDoc).filter(KnowledgeDoc.title == f"单独{mark}").one()
        assert alone.status == "active"
    finally:
        app.dependency_overrides.pop(get_current_agent, None)
        _cleanup(db, titles, agent.id)


def test_create_faq_rejects_duplicate_title(db):
    mark = uuid.uuid4().hex[:8]
    title = f"重建标题{mark}"
    agent = _agent(db)
    app.dependency_overrides[get_current_agent] = lambda: agent
    client = TestClient(app)
    payload = {"title": title, "question": "怎么做", "answer": "就这样"}
    try:
        created = client.post("/api/knowledge/faq", json=payload)
        assert created.status_code == 200, created.text
        again = client.post("/api/knowledge/faq", json=payload)
        assert again.status_code == 409, again.text
        assert "该标题已存在" in again.json()["detail"]
        db.expire_all()
        assert db.query(KnowledgeDoc).filter(
            KnowledgeDoc.title == title, KnowledgeDoc.status != "archived",
        ).count() == 1
    finally:
        app.dependency_overrides.pop(get_current_agent, None)
        _cleanup(db, [title], agent.id)
